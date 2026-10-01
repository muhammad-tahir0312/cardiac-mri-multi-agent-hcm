from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt


SOURCE = Path("ethics/umrec/UMREC_Form_A_Exempt_Research_Review_2025.docx")
OUTPUT = Path("ethics/umrec/UMREC_Form_A_DRAFT_Muhammad_Tahir.docx")


PROJECT_TITLE = (
    "A Multi-Agent System for Automated Cardiac MRI Analysis and "
    "Guideline-Grounded Decision Support"
)

SUMMARY = (
    "This Master of Computer Science by Research project evaluates a multi-agent cardiac "
    "magnetic resonance (CMR) pipeline that integrates automated short-axis cine segmentation, "
    "deterministic ventricular quantification, guideline-grounded report generation and "
    "failure-aware quality control. The primary objective is to determine whether the pipeline "
    "generalises to an independent external cohort and whether segmentation errors propagate "
    "into clinically relevant measures such as left- and right-ventricular end-diastolic volume, "
    "end-systolic volume, ejection fraction and myocardial mass. Secondary analyses will assess "
    "failure detection and, only where suitable de-identified labels or reports are available, "
    "the accuracy and grounding of generated reports.\n\n"
    "The study will use an existing dataset supplied under controlled access by an external "
    "custodian such as the SCMR Registry or the BAAI Cardiac Agent research team. No participants "
    "will be recruited or contacted, and no new scans will be acquired. No UMMC patient data will "
    "be used. The requested data are de-identified short-axis cine CMR images, end-diastolic and "
    "end-systolic frames or complete cine series, reference contours or derived ventricular "
    "measurements, acquisition metadata, pseudonymous patient identifiers and diagnostic labels "
    "where permitted. The final custodian and approved variables will be recorded before analysis; "
    "written access permission and any data-use or material-transfer agreement will be obtained.\n\n"
    "The custodian must remove direct identifiers before access. The research team will not "
    "attempt re-identification or linkage to external identifiable records. Only the minimum "
    "approved variables will be received, stored in an access-controlled Universiti Malaya "
    "environment and accessed by named researchers. Data will not be placed in personal email, "
    "consumer cloud storage, public repositories or unapproved AI services. Results will be "
    "reported only in aggregate, and small subgroups will be suppressed where disclosure risk exists.\n\n"
    "The external cohort will be held as a locked test set and will not be used to tune thresholds "
    "or train the model. Segmentation will be evaluated using patient-level Dice, surface-distance "
    "and failure-rate measures. Ventricular measurements will be compared with available reference "
    "values using absolute error, bias, limits of agreement and clinically defined tolerance "
    "thresholds. Performance will be stratified by scanner vendor, site and diagnosis when sample "
    "sizes permit. All exclusions, missing data and analysis decisions will be documented before "
    "outcome inspection."
)


def set_cell(cell, text, *, highlight=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.font.size = Pt(9)
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


doc = Document(SOURCE)

# Project title.
set_cell(doc.tables[1].cell(0, 1), PROJECT_TITLE)

# Staff principal investigator / supervisor.
pi_row = doc.tables[3].rows[1].cells
set_cell(pi_row[0], "Dr")
set_cell(pi_row[1], "Uzair Iqbal")
set_cell(pi_row[2], "03-7967 6306")
set_cell(pi_row[3], "uzairiqbal@um.edu.my")

# Other investigators require confirmation because the thesis records a co-supervisor,
# while the SCMR feasibility request was submitted with no co-investigator.
other_row = doc.tables[5].rows[1].cells
set_cell(other_row[0], "TBC", highlight=True)
set_cell(other_row[1], "Confirm co-supervisor listing", highlight=True)
set_cell(other_row[2], "TBC", highlight=True)
set_cell(other_row[3], "TBC", highlight=True)

# Student principal investigator.
student_row = doc.tables[7].rows[1].cells
set_cell(student_row[0], "Mr")
set_cell(student_row[1], "Muhammad Tahir")
set_cell(student_row[2], "TO CONFIRM", highlight=True)
set_cell(student_row[3], "mtahirkhanyousufzai@gmail.com")

# The supplied form has no Master by Research checkbox. State the actual programme and
# flag the category rather than selecting an inaccurate option.
degree_paragraph = doc.paragraphs[53]
degree_paragraph.clear()
degree_paragraph.add_run("2.4\tDEGREE/PROGRAMME:\t").bold = True
degree_note = degree_paragraph.add_run("Master (Research) - confirm UMREC checkbox.")
degree_note.font.highlight_color = WD_COLOR_INDEX.YELLOW

# Grant information and study dates are not present in the supplied research files.
set_cell(doc.tables[10].cell(0, 1), "TO CONFIRM", highlight=True)
set_cell(doc.tables[10].cell(1, 1), "TO CONFIRM", highlight=True)
set_cell(doc.tables[11].cell(0, 3), "TO CONFIRM", highlight=True)
set_cell(doc.tables[11].cell(0, 5), "TO CONFIRM", highlight=True)
set_cell(doc.tables[12].cell(0, 3), "TO CONFIRM", highlight=True)
set_cell(doc.tables[12].cell(0, 5), "TO CONFIRM", highlight=True)

# Insert the required 300-400 word plain-English project summary.
summary_paragraph = doc.paragraphs[68]
summary_paragraph.text = SUMMARY
summary_paragraph.paragraph_format.space_after = Pt(6)
for run in summary_paragraph.runs:
    run.font.size = Pt(10)

# Remove the template's unused blank answer-line paragraphs after the inserted
# summary. Leaving them in place creates an unnecessary blank page.
for paragraph in reversed(doc.paragraphs[69:85]):
    delete_paragraph(paragraph)

# Names can be prepared, but signatures and dates must be completed by the people concerned.
set_cell(doc.tables[13].cell(0, 0), "Dr Uzair Iqbal")
set_cell(doc.tables[13].cell(0, 1), "SIGNATURE REQUIRED", highlight=True)
set_cell(doc.tables[13].cell(0, 3), "DATE REQUIRED", highlight=True)
set_cell(doc.tables[14].cell(0, 0), "Dr Uzair Iqbal")
set_cell(doc.tables[14].cell(0, 1), "SIGNATURE REQUIRED", highlight=True)
set_cell(doc.tables[14].cell(0, 3), "DATE REQUIRED", highlight=True)
set_cell(doc.tables[15].cell(0, 1), "TO BE COMPLETED BY DEPUTY DEAN", highlight=True)
set_cell(doc.tables[15].cell(0, 3), "DATE REQUIRED", highlight=True)
set_cell(
    doc.tables[15].cell(2, 1),
    "Faculty of Computer Science and Information Technology",
)

doc.save(OUTPUT)
print(f"Wrote {OUTPUT}")
print(f"Summary words: {len(SUMMARY.split())}")
